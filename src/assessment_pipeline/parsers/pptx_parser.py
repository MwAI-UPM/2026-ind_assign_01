"""
PowerPoint parsing utilities.
"""

from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation


LOGGER = logging.getLogger(__name__)


class PptxParser:
    """
    Extract textual content from PowerPoint presentations.
    """

    def extract_text(
        self,
        pptx_path: Path,
    ) -> str:
        """
        Extract textual content from a PPTX presentation.

        Args:
            pptx_path:
                Path to the presentation.

        Returns:
            Extracted textual content.
        """

        if not pptx_path.exists():
            raise FileNotFoundError(
                f"Presentation not found: {pptx_path}"
            )

        LOGGER.info(
            "Parsing PPTX: %s",
            pptx_path,
        )

        presentation = Presentation(
            str(pptx_path)
        )

        fragments: list[str] = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            fragments.append(
                f"=== SLIDE {slide_number} ==="
            )

            for shape in slide.shapes:

                if not hasattr(
                    shape,
                    "text",
                ):
                    continue

                text = shape.text.strip()

                if text:
                    fragments.append(text)

        return "\n".join(fragments)
