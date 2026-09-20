"""
Excel parsing utilities.
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook


LOGGER = logging.getLogger(__name__)


class XlsxParser:
    """
    Extract textual content from Microsoft Excel workbooks.
    """

    def extract_text(
        self,
        xlsx_path: Path,
    ) -> str:
        """
        Extract textual content from an XLSX workbook.

        Args:
            xlsx_path:
                Path to the workbook.

        Returns:
            Extracted textual representation.
        """

        if not xlsx_path.exists():
            raise FileNotFoundError(
                f"Workbook not found: {xlsx_path}"
            )

        LOGGER.info(
            "Parsing XLSX: %s",
            xlsx_path,
        )

        workbook = load_workbook(
            filename=xlsx_path,
            read_only=True,
            data_only=True,
        )

        fragments: list[str] = []

        for sheet in workbook.worksheets:

            fragments.append(
                f"=== SHEET: {sheet.title} ==="
            )

            for row in sheet.iter_rows():

                values = []

                for cell in row:

                    if cell.value is None:
                        continue

                    values.append(
                        str(cell.value)
                    )

                if values:
                    fragments.append(
                        " | ".join(values)
                    )

        return "\n".join(fragments)
