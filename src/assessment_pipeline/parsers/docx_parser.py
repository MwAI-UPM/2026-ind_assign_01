"""
DOCX parsing utilities.
"""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document


LOGGER = logging.getLogger(__name__)


class DocxParser:
    """
    Extract textual content from Microsoft Word documents.
    """

    def extract_text(
        self,
        docx_path: Path,
    ) -> str:
        """
        Extract textual content from a DOCX file.

        Args:
            docx_path:
                Path to the DOCX file.

        Returns:
            Extracted textual content.

        Raises:
            FileNotFoundError:
                If the document does not exist.
        """

        if not docx_path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {docx_path}"
            )

        LOGGER.info(
            "Parsing DOCX: %s",
            docx_path,
        )

        document = Document(docx_path)

        fragments: list[str] = []

        #
        # Paragraphs
        #
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                fragments.append(text)

        #
        # Tables
        #
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    fragments.append(
                        " | ".join(cells)
                    )

        return "\n".join(fragments)
